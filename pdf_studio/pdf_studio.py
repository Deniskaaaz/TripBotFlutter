"""
PDF Студия — конвертер и редактор PDF
Зависимости: pip install PySide6 PyMuPDF python-docx python-pptx
Опционально: LibreOffice (для качественной конвертации DOCX/PPTX/XLSX)
Запуск: py pdf_studio.py
"""

import sys
import os
import shutil
import subprocess
import tempfile
import pymupdf as fitz  # PyMuPDF

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QFileDialog, QGraphicsView, QGraphicsScene,
    QGraphicsPixmapItem, QGraphicsTextItem, QGraphicsPathItem,
    QGraphicsRectItem, QScrollArea,
    QMessageBox, QInputDialog, QDialog, QColorDialog,
    QSlider, QFrame, QButtonGroup, QGraphicsItem, QStatusBar,
    QTabWidget, QListWidget, QListWidgetItem, QSizePolicy
)
from PySide6.QtGui import (
    QPixmap, QImage, QPainter, QPen, QBrush, QColor, QFont, QAction,
    QPainterPath, QPainterPathStroker, QKeySequence, QIcon, QPalette
)
from PySide6.QtCore import (
    Qt, QBuffer, QIODevice, QSize, QDateTime, QRectF, QSettings
)

# ─────────────────────────────────────────────────────────────
#  Константы и пути
# ─────────────────────────────────────────────────────────────
ZOOM = 2.0
PDF_SCALE = 1.0 / ZOOM

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SIGNATURES_DIR = os.path.join(BASE_DIR, "signatures")
os.makedirs(SIGNATURES_DIR, exist_ok=True)

# UI-константы
UI_PANEL_WIDTH = 260
UI_TOOL_H = 40
UI_ACTION_H = 36
UI_NAV_H = 34
UI_BASE_FONT = 10

# Сдвиг при редактировании
LINE_HEIGHT_FACTOR = 1.25     # межстрочный интервал (в единицах размера шрифта)


def get_font_path():
    candidates = [
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/calibri.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
        "/Library/Fonts/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/TTF/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    ]
    for p in candidates:
        if os.path.exists(p):
            return p
    return None


FONT_PATH = get_font_path()

_FONT_VARIANTS = {
    ('arial',     True,  False): 'arialbd.ttf',
    ('arial',     False, True):  'ariali.ttf',
    ('arial',     True,  True):  'arialbi.ttf',
    ('segoeui',   True,  False): 'segoeuib.ttf',
    ('segoeui',   False, True):  'segoeuii.ttf',
    ('segoeui',   True,  True):  'segoeuiz.ttf',
    ('calibri',   True,  False): 'calibrib.ttf',
    ('calibri',   False, True):  'calibrii.ttf',
    ('calibri',   True,  True):  'calibriz.ttf',
    ('dejavusans', True, False): 'DejaVuSans-Bold.ttf',
    ('dejavusans', False, True): 'DejaVuSans-Oblique.ttf',
    ('dejavusans', True, True):  'DejaVuSans-BoldOblique.ttf',
    ('liberationsans', True, False): 'LiberationSans-Bold.ttf',
    ('liberationsans', False, True): 'LiberationSans-Italic.ttf',
    ('liberationsans', True, True):  'LiberationSans-BoldItalic.ttf',
}


def find_font_variant(base_path, bold, italic):
    if not base_path or (not bold and not italic):
        return None
    folder = os.path.dirname(base_path)
    name, _ = os.path.splitext(os.path.basename(base_path))
    fname = _FONT_VARIANTS.get((name.lower(), bold, italic))
    if fname:
        candidate = os.path.join(folder, fname)
        if os.path.exists(candidate):
            return candidate
    return None


def pixmap_to_qimage(pix):
    fmt = QImage.Format_RGBA8888 if pix.alpha else QImage.Format_RGB888
    img = QImage(pix.samples, pix.width, pix.height, pix.stride, fmt)
    return img.copy()


# ─────────────────────────────────────────────────────────────
#  LibreOffice
# ─────────────────────────────────────────────────────────────
def find_libreoffice():
    candidates = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice",
        "/usr/local/bin/soffice",
        "/opt/libreoffice/program/soffice",
    ]
    for p in candidates:
        if os.path.exists(p):
            return p
    return shutil.which("soffice") or shutil.which("libreoffice")


LIBREOFFICE = find_libreoffice()


def libreoffice_convert(src_path):
    if not LIBREOFFICE:
        return None

    tmpdir = tempfile.mkdtemp(prefix="pdfstudio_")
    profile_dir = tempfile.mkdtemp(prefix="lo_profile_")
    profile_url = "file:///" + profile_dir.replace("\\", "/")

    try:
        cmd = [
            LIBREOFFICE,
            "--headless", "--norestore", "--nologo",
            f"-env:UserInstallation={profile_url}",
            "--convert-to", "pdf",
            "--outdir", tmpdir,
            src_path,
        ]
        flags = 0x08000000 if sys.platform.startswith("win") else 0
        subprocess.run(cmd, capture_output=True, timeout=180,
                       creationflags=flags)

        base = os.path.splitext(os.path.basename(src_path))[0] + ".pdf"
        out = os.path.join(tmpdir, base)
        if os.path.exists(out):
            return out
    except Exception as e:
        print("LibreOffice error:", e)

    try:
        shutil.rmtree(tmpdir, ignore_errors=True)
        shutil.rmtree(profile_dir, ignore_errors=True)
    except Exception:
        pass
    return None


def docx_to_pdf_fallback(path, out_doc):
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError(
            "Для DOCX без LibreOffice нужен python-docx:\n"
            "py -m pip install python-docx"
        )

    d = Document(path)
    A4 = (595, 842)
    margin = 50
    page = out_doc.new_page(width=A4[0], height=A4[1])
    y = margin
    line_h = 16

    def new_page_if_needed():
        nonlocal page, y
        if y > A4[1] - margin:
            page = out_doc.new_page(width=A4[0], height=A4[1])
            y = margin

    kwargs = {}
    if FONT_PATH:
        kwargs['fontfile'] = FONT_PATH
        kwargs['fontname'] = "custom"
    else:
        kwargs['fontname'] = "helv"

    for para in d.paragraphs:
        text = para.text
        if not text.strip():
            y += line_h // 2
            continue
        new_page_if_needed()
        try:
            page.insert_text(
                fitz.Point(margin, y + 12), text,
                fontsize=12, color=(0, 0, 0), **kwargs
            )
        except Exception:
            pass
        y += line_h

    for table in d.tables:
        new_page_if_needed()
        for row in table.rows:
            cells = " | ".join(c.text.strip() for c in row.cells)
            if not cells.strip():
                continue
            new_page_if_needed()
            try:
                page.insert_text(
                    fitz.Point(margin, y + 12), cells,
                    fontsize=11, color=(0, 0, 0), **kwargs
                )
            except Exception:
                pass
            y += line_h


def pptx_to_pdf_fallback(path, out_doc):
    try:
        from pptx import Presentation
    except ImportError:
        raise RuntimeError(
            "Для PPTX без LibreOffice нужен python-pptx:\n"
            "py -m pip install python-pptx"
        )

    prs = Presentation(path)
    W, H = 842, 595
    margin = 50

    kwargs = {}
    if FONT_PATH:
        kwargs['fontfile'] = FONT_PATH
        kwargs['fontname'] = "custom"
    else:
        kwargs['fontname'] = "helv"

    for idx, slide in enumerate(prs.slides, 1):
        page = out_doc.new_page(width=W, height=H)
        y = margin
        try:
            page.insert_text(
                fitz.Point(margin, y + 14), f"— Слайд {idx} —",
                fontsize=10, color=(0.4, 0.4, 0.4), **kwargs
            )
        except Exception:
            pass
        y += 28

        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if not text.strip():
                        y += 8
                        continue
                    if y > H - margin:
                        page = out_doc.new_page(width=W, height=H)
                        y = margin
                    try:
                        page.insert_text(
                            fitz.Point(margin, y + 14), text,
                            fontsize=14, color=(0, 0, 0), **kwargs
                        )
                    except Exception:
                        pass
                    y += 20


# ─────────────────────────────────────────────────────────────
#  Расширенная зона клика
# ─────────────────────────────────────────────────────────────
HIT_PADDING = 6
HIT_MIN_WIDTH = 14


class SelectableTextItem(QGraphicsTextItem):
    def shape(self):
        p = QPainterPath()
        p.addRect(self.boundingRect().adjusted(
            -HIT_PADDING, -HIT_PADDING, HIT_PADDING, HIT_PADDING))
        return p


class SelectablePixmapItem(QGraphicsPixmapItem):
    def shape(self):
        p = QPainterPath()
        p.addRect(self.boundingRect().adjusted(
            -HIT_PADDING, -HIT_PADDING, HIT_PADDING, HIT_PADDING))
        return p


class SelectablePathItem(QGraphicsPathItem):
    def shape(self):
        stroker = QPainterPathStroker()
        w = max(self.pen().widthF() + HIT_PADDING * 2, HIT_MIN_WIDTH)
        stroker.setWidth(w)
        return stroker.createStroke(self.path())


class SelectableRectItem(QGraphicsRectItem):
    def shape(self):
        p = QPainterPath()
        p.addRect(self.rect())
        return p


# ─────────────────────────────────────────────────────────────
#  Диалог подписи
# ─────────────────────────────────────────────────────────────
class SignatureDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Подпись")
        self.resize(660, 480)

        self.result_pixmap = None
        layout = QVBoxLayout(self)

        self.tabs = QTabWidget()
        layout.addWidget(self.tabs)

        self.tab_draw = QWidget()
        tdl = QVBoxLayout(self.tab_draw)

        self.canvas = QPixmap(600, 240)
        self.canvas.fill(Qt.transparent)

        self.label = QLabel()
        self.label.setPixmap(self.canvas)
        self.label.setFixedSize(600, 240)
        self.label.setStyleSheet("border: 1px solid #888; background: white;")
        self.label.setAlignment(Qt.AlignCenter)
        tdl.addWidget(self.label)

        draw_btns = QHBoxLayout()
        clear_btn = QPushButton("Очистить")
        clear_btn.setMinimumHeight(UI_ACTION_H)
        clear_btn.clicked.connect(self.clear_canvas)
        save_btn = QPushButton("💾 Сохранить в библиотеку")
        save_btn.setMinimumHeight(UI_ACTION_H)
        save_btn.clicked.connect(self.save_to_library)
        draw_btns.addWidget(clear_btn)
        draw_btns.addStretch()
        draw_btns.addWidget(save_btn)
        tdl.addLayout(draw_btns)

        self.tabs.addTab(self.tab_draw, "✏  Нарисовать")

        self.tab_saved = QWidget()
        tsl = QVBoxLayout(self.tab_saved)

        self.list_widget = QListWidget()
        self.list_widget.setIconSize(QSize(260, 90))
        self.list_widget.itemDoubleClicked.connect(self._use_saved_double)
        tsl.addWidget(self.list_widget)

        saved_btns = QHBoxLayout()
        for text, slot in [
            ("✓ Использовать", self.use_saved),
            ("🗑 Удалить", self.delete_saved),
        ]:
            b = QPushButton(text)
            b.setMinimumHeight(UI_ACTION_H)
            b.clicked.connect(slot)
            saved_btns.addWidget(b)
        saved_btns.addStretch()
        refresh_btn = QPushButton("🔄 Обновить")
        refresh_btn.setMinimumHeight(UI_ACTION_H)
        refresh_btn.clicked.connect(self.load_saved_list)
        saved_btns.addWidget(refresh_btn)
        tsl.addLayout(saved_btns)

        self.tabs.addTab(self.tab_saved, "📚  Сохранённые")

        bottom = QHBoxLayout()
        ok_btn = QPushButton("Вставить")
        ok_btn.setMinimumHeight(UI_ACTION_H)
        ok_btn.clicked.connect(self._accept_draw)
        cancel_btn = QPushButton("Отмена")
        cancel_btn.setMinimumHeight(UI_ACTION_H)
        cancel_btn.clicked.connect(self.reject)
        bottom.addStretch()
        bottom.addWidget(cancel_btn)
        bottom.addWidget(ok_btn)
        layout.addLayout(bottom)

        self.drawing = False
        self.last_pos = None
        self.label.mousePressEvent = self.on_press
        self.label.mouseMoveEvent = self.on_move
        self.label.mouseReleaseEvent = self.on_release

        self.load_saved_list()
        if self.list_widget.count() > 0:
            self.tabs.setCurrentIndex(1)

    def clear_canvas(self):
        self.canvas.fill(Qt.transparent)
        self.label.setPixmap(self.canvas)

    def on_press(self, e):
        self.drawing = True
        self.last_pos = e.pos()

    def on_move(self, e):
        if not self.drawing:
            return
        p = QPainter(self.canvas)
        p.setRenderHint(QPainter.Antialiasing)
        pen = QPen(Qt.black, 3, Qt.SolidLine, Qt.RoundCap, Qt.RoundJoin)
        p.setPen(pen)
        p.drawLine(self.last_pos, e.pos())
        p.end()
        self.label.setPixmap(self.canvas)
        self.last_pos = e.pos()

    def on_release(self, e):
        self.drawing = False

    def _trimmed_pixmap(self):
        img = self.canvas.toImage().convertToFormat(QImage.Format_ARGB32)
        w, h = img.width(), img.height()
        min_x, min_y = w, h
        max_x, max_y = -1, -1
        bits = img.constBits()
        bpl = img.bytesPerLine()
        for y in range(h):
            row = y * bpl
            for x in range(w):
                if bits[row + x * 4 + 3] > 10:
                    if x < min_x: min_x = x
                    if y < min_y: min_y = y
                    if x > max_x: max_x = x
                    if y > max_y: max_y = y
        if max_x < 0:
            return None
        pad = 6
        min_x = max(0, min_x - pad)
        min_y = max(0, min_y - pad)
        max_x = min(w - 1, max_x + pad)
        max_y = min(h - 1, max_y + pad)
        cw = max_x - min_x + 1
        ch = max_y - min_y + 1
        return QPixmap.fromImage(img.copy(min_x, min_y, cw, ch))

    def _accept_draw(self):
        pm = self._trimmed_pixmap()
        if pm is None:
            QMessageBox.warning(self, "Пусто", "Сначала нарисуйте подпись.")
            return
        self.result_pixmap = pm
        self.accept()

    def save_to_library(self):
        pm = self._trimmed_pixmap()
        if pm is None:
            QMessageBox.warning(self, "Пусто", "Нечего сохранять — нарисуйте подпись.")
            return

        name, ok = QInputDialog.getText(
            self, "Сохранить подпись",
            "Название (например, «Моя подпись»):",
            text="Подпись_" + QDateTime.currentDateTime().toString("ddMMyyyy_hhmmss"))
        if not ok or not name.strip():
            return

        safe = "".join(c for c in name if c not in r'\/:*?"<>|').strip()
        if not safe:
            safe = "signature"
        path = os.path.join(SIGNATURES_DIR, safe + ".png")

        if os.path.exists(path):
            if QMessageBox.question(self, "Перезаписать?",
                                    f"«{safe}.png» уже существует. Заменить?"
                                    ) != QMessageBox.Yes:
                return
        if not pm.save(path, "PNG"):
            QMessageBox.critical(self, "Ошибка", "Не удалось сохранить файл.")
            return

        QMessageBox.information(self, "Готово", f"Сохранено:\n{path}")
        self.load_saved_list()
        self.tabs.setCurrentIndex(1)

    def load_saved_list(self):
        self.list_widget.clear()
        if not os.path.isdir(SIGNATURES_DIR):
            return
        files = sorted(f for f in os.listdir(SIGNATURES_DIR)
                       if f.lower().endswith(".png"))
        for f in files:
            full = os.path.join(SIGNATURES_DIR, f)
            pm = QPixmap(full)
            if pm.isNull():
                continue
            thumb = pm.scaled(260, 90, Qt.KeepAspectRatio, Qt.SmoothTransformation)
            item = QListWidgetItem(QIcon(thumb), os.path.splitext(f)[0])
            item.setData(Qt.UserRole, full)
            self.list_widget.addItem(item)

    def _use_saved_double(self, item):
        self.use_saved()

    def use_saved(self):
        item = self.list_widget.currentItem()
        if item is None:
            QMessageBox.information(self, "Не выбрано",
                                    "Выберите подпись из списка.")
            return
        path = item.data(Qt.UserRole)
        pm = QPixmap(path)
        if pm.isNull():
            QMessageBox.warning(self, "Ошибка", "Файл подписи повреждён.")
            return
        self.result_pixmap = pm
        self.accept()

    def delete_saved(self):
        item = self.list_widget.currentItem()
        if item is None:
            return
        name = item.text()
        if QMessageBox.question(self, "Удалить",
                                f"Удалить подпись «{name}»?"
                                ) != QMessageBox.Yes:
            return
        path = item.data(Qt.UserRole)
        try:
            os.remove(path)
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", str(e))
            return
        self.load_saved_list()


# ─────────────────────────────────────────────────────────────
#  View
# ─────────────────────────────────────────────────────────────
class PageView(QGraphicsView):
    def __init__(self, main_window):
        super().__init__()
        self.main = main_window
        self.setRenderHint(QPainter.Antialiasing)
        self.setDragMode(QGraphicsView.RubberBandDrag)
        self.setBackgroundBrush(QBrush(QColor(50, 50, 55)))
        self.setMouseTracking(True)
        self.setAlignment(Qt.AlignCenter)

        self.drawing = False
        self.draw_path = None
        self.draw_item = None

        self.moving = False
        self.move_start = None
        self.move_items = []
        self.move_orig_pos = []

    def _is_movable(self, item):
        return (item is not None
                and item.zValue() >= 0
                and bool(item.flags() & QGraphicsItem.ItemIsMovable))

    def _update_hover_cursor(self, view_pos):
        if self.main.current_tool != 'select':
            return
        item = self.itemAt(view_pos)
        if self._is_movable(item):
            self.viewport().setCursor(Qt.OpenHandCursor)
        else:
            self.viewport().setCursor(Qt.ArrowCursor)

    def leaveEvent(self, event):
        if not self.moving and self.main.current_tool == 'select':
            self.viewport().setCursor(Qt.ArrowCursor)
        super().leaveEvent(event)

    def mousePressEvent(self, event):
        if event.button() != Qt.LeftButton or self.main.pdf is None:
            super().mousePressEvent(event)
            return

        tool = self.main.current_tool
        pos = self.mapToScene(event.pos())

        if tool == 'select':
            self._handle_select_press(event)
        elif tool == 'text':
            self.main.add_text_at(pos)
        elif tool == 'image':
            self.main.add_image_at(pos)
        elif tool == 'signature':
            self.main.add_signature_at(pos)
        elif tool == 'edit_existing':
            self.main.edit_existing_text_at(pos)
        elif tool in ('draw', 'whiteout'):
            self._start_draw(pos, tool)
        else:
            super().mousePressEvent(event)

    def _handle_select_press(self, event):
        scene_pos = self.mapToScene(event.pos())
        item = self.itemAt(event.pos())

        if self._is_movable(item):
            ctrl = bool(event.modifiers() & Qt.ControlModifier)
            if not item.isSelected():
                if not ctrl:
                    self.scene().clearSelection()
                item.setSelected(True)

            self.moving = True
            self.move_start = scene_pos
            self.move_items = list(self.scene().selectedItems())
            self.move_orig_pos = [it.pos() for it in self.move_items]
            self.viewport().setCursor(Qt.ClosedHandCursor)
            event.accept()
        else:
            if not (event.modifiers() & Qt.ControlModifier):
                self.scene().clearSelection()
            self.moving = False
            super().mousePressEvent(event)

    def mouseMoveEvent(self, event):
        if self.moving:
            delta = self.mapToScene(event.pos()) - self.move_start
            for it, orig in zip(self.move_items, self.move_orig_pos):
                it.setPos(orig + delta)
            event.accept()
            return

        if self.drawing:
            pos = self.mapToScene(event.pos())
            self.draw_path.lineTo(pos)
            self.draw_item.setPath(self.draw_path)
            event.accept()
            return

        self._update_hover_cursor(event.pos())
        super().mouseMoveEvent(event)

    def mouseReleaseEvent(self, event):
        if self.moving:
            self.moving = False
            self.move_items = []
            self.move_orig_pos = []
            self._update_hover_cursor(event.pos())
            event.accept()
            return

        if self.drawing:
            self.drawing = False
            self.main.register_annotation('drawing', self.draw_item)
            self.draw_item = None
            self.draw_path = None
            event.accept()
            return

        super().mouseReleaseEvent(event)

    def _start_draw(self, pos, tool):
        self.drawing = True
        self.draw_path = QPainterPath(pos)

        is_white = (tool == 'whiteout')
        color = QColor("#FFFFFF") if is_white else self.main.current_color
        width = 20 if is_white else self.main.current_width

        pen = QPen(color, width, Qt.SolidLine, Qt.RoundCap, Qt.RoundJoin)
        self.draw_item = SelectablePathItem()
        self.draw_item.setPen(pen)
        self.draw_item.setBrush(Qt.NoBrush)
        self.draw_item.setZValue(10)
        self.draw_item.setFlag(QGraphicsItem.ItemIsSelectable)
        self.draw_item.setData(0, {
            'kind': 'whiteout' if is_white else 'draw',
            'color': color.name(),
            'width': width,
        })
        self.scene().addItem(self.draw_item)
        self.draw_item.setPath(self.draw_path)


# ─────────────────────────────────────────────────────────────
#  Главное окно
# ─────────────────────────────────────────────────────────────
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("PDF Студия")
        self.setMinimumSize(1000, 640)

        self.pdf = None
        self.current_page = 0
        self.current_tool = 'select'
        self.current_color = QColor(0, 0, 0)
        self.current_width = 3

        self.annotations = {}
        self.page_pixmap_item = None

        self.settings = QSettings("PDFStudio", "MainWindow")

        self._build_ui()
        self._setup_shortcuts()
        self._restore_window_state()

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root = QHBoxLayout(central)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        panel_container = QWidget()
        panel_container.setFixedWidth(UI_PANEL_WIDTH)
        panel_container.setStyleSheet("QWidget { background: #2b2b33; }")

        outer = QVBoxLayout(panel_container)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        scroll.setStyleSheet("""
            QScrollArea { background: #2b2b33; border: none; }
            QScrollBar:vertical {
                background: #2b2b33; width: 8px; margin: 0;
            }
            QScrollBar::handle:vertical {
                background: #4a4a58; border-radius: 4px; min-height: 30px;
            }
            QScrollBar::handle:vertical:hover { background: #5a5a68; }
            QScrollBar::add-line:vertical,
            QScrollBar::sub-line:vertical { height: 0; }
            QScrollBar::add-page:vertical,
            QScrollBar::sub-page:vertical { background: transparent; }
        """)
        outer.addWidget(scroll)

        panel_inner = QWidget()
        panel_inner.setStyleSheet("QWidget { background: #2b2b33; }")
        scroll.setWidget(panel_inner)

        pl = QVBoxLayout(panel_inner)
        pl.setContentsMargins(10, 12, 10, 12)
        pl.setSpacing(6)

        btn_style = """
            QPushButton {
                padding: 8px 12px;
                background: #383843; color: #e0e0e0;
                border: none; border-radius: 6px;
                text-align: left;
                font-size: 13px;
            }
            QPushButton:hover { background: #4a4a58; }
            QPushButton:checked { background: #d04848; color: white; }
        """

        hdr = QLabel("ИНСТРУМЕНТЫ")
        hdr.setStyleSheet(
            "color:#ff6b6b; font-weight:bold; font-size:12px; padding:4px 4px 8px 4px;"
        )
        pl.addWidget(hdr)

        self.tool_buttons = {}
        self.tool_group = QButtonGroup(self)
        self.tool_group.setExclusive(True)
        for tool, label in [
            ('select',        '↖  Выбор'),
            ('text',          'T  Текст'),
            ('image',         '🖼  Изображение'),
            ('signature',     '✎  Подпись'),
            ('draw',          '✏  Рисование'),
            ('whiteout',      '▭  Замазка'),
            ('edit_existing', '📝  Изменить текст'),
        ]:
            b = QPushButton(label)
            b.setCheckable(True)
            b.setFixedHeight(UI_TOOL_H)
            b.setStyleSheet(btn_style)
            if tool == 'select':
                b.setChecked(True)
            b.clicked.connect(lambda _, t=tool: self.set_tool(t))
            pl.addWidget(b)
            self.tool_group.addButton(b)
            self.tool_buttons[tool] = b

        hint = QLabel(
            "📝 кликните по тексту в PDF,\n"
            "чтобы изменить его.\n"
            "Всё, что ниже, аккуратно\n"
            "сдвинется, если строк\n"
            "станет больше."
        )
        hint.setStyleSheet("color:#888; font-size:11px; padding:2px 6px 6px 6px;")
        hint.setWordWrap(True)
        pl.addWidget(hint)

        def sep():
            line = QFrame()
            line.setFrameShape(QFrame.HLine)
            line.setStyleSheet(
                "color: #3a3a45; background: #3a3a45; max-height: 1px;"
            )
            pl.addSpacing(4)
            pl.addWidget(line)
            pl.addSpacing(4)

        sep()

        lbl_color = QLabel("Цвет:")
        lbl_color.setStyleSheet("color:#cfcfcf; font-size:12px;")
        pl.addWidget(lbl_color)

        self.color_btn = QPushButton()
        self.color_btn.setFixedHeight(UI_ACTION_H)
        self.color_btn.setStyleSheet(btn_style)
        self._update_color_btn()
        self.color_btn.clicked.connect(self._pick_color)
        pl.addWidget(self.color_btn)

        lbl_size = QLabel("Толщина:")
        lbl_size.setStyleSheet("color:#cfcfcf; font-size:12px; padding-top:4px;")
        pl.addWidget(lbl_size)

        self.width_slider = QSlider(Qt.Horizontal)
        self.width_slider.setRange(1, 30)
        self.width_slider.setValue(3)
        self.width_slider.setFixedHeight(24)
        self.width_slider.valueChanged.connect(self._set_width)
        pl.addWidget(self.width_slider)

        self.width_lbl = QLabel("3 px")
        self.width_lbl.setStyleSheet("color:#cfcfcf; font-size:12px;")
        pl.addWidget(self.width_lbl)

        sep()

        for text, slot in [
            ("📂  Открыть PDF",     self.open_pdf),
            ("📄  Новый документ",  self.new_document),
            ("🔄  Конвертировать…", self.convert_files),
            ("💾  Сохранить PDF…",  self.export_pdf),
            ("📁  Папка подписей",  self.open_signatures_folder),
        ]:
            b = QPushButton(text)
            b.setFixedHeight(UI_ACTION_H)
            b.setStyleSheet(btn_style)
            b.clicked.connect(slot)
            pl.addWidget(b)

        pl.addStretch(1)

        sep()

        nav = QHBoxLayout()
        nav.setSpacing(6)
        prev = QPushButton("◀")
        nxt = QPushButton("▶")
        prev.setFixedSize(44, UI_NAV_H)
        nxt.setFixedSize(44, UI_NAV_H)
        prev.setStyleSheet(btn_style + "QPushButton{ text-align:center; }")
        nxt.setStyleSheet(btn_style + "QPushButton{ text-align:center; }")
        prev.clicked.connect(lambda: self.change_page(-1))
        nxt.clicked.connect(lambda: self.change_page(1))

        self.page_lbl = QLabel("— / —")
        self.page_lbl.setAlignment(Qt.AlignCenter)
        self.page_lbl.setStyleSheet(
            "color:#e0e0e0; font-size:13px; font-weight:bold;"
        )
        nav.addWidget(prev)
        nav.addWidget(self.page_lbl, 1)
        nav.addWidget(nxt)
        pl.addLayout(nav)

        root.addWidget(panel_container)

        self.scene = QGraphicsScene()
        self.view = PageView(self)
        self.view.setScene(self.scene)
        self.view.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        root.addWidget(self.view, 1)

        self.setStatusBar(QStatusBar())
        self.statusBar().showMessage(
            "Готово. Откройте PDF или конвертируйте файлы."
        )

    def _setup_shortcuts(self):
        a = QAction(self)
        a.setShortcut(QKeySequence.Delete)
        a.triggered.connect(self.delete_selected)
        self.addAction(a)

        esc = QAction(self)
        esc.setShortcut(QKeySequence(Qt.Key_Escape))
        esc.triggered.connect(lambda: self.set_tool('select'))
        self.addAction(esc)

    def _restore_window_state(self):
        geom = self.settings.value("geometry")
        state = self.settings.value("windowState")
        if geom is not None:
            self.restoreGeometry(geom)
        if state is not None:
            self.restoreState(state)

    def closeEvent(self, event):
        self.settings.setValue("geometry", self.saveGeometry())
        self.settings.setValue("windowState", self.saveState())
        super().closeEvent(event)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if self.pdf and self.page_pixmap_item is not None:
            self.view.fitInView(self.scene.sceneRect(), Qt.KeepAspectRatio)

    # ─── Инструменты ───
    def set_tool(self, tool):
        self.current_tool = tool
        if tool in self.tool_buttons:
            self.tool_buttons[tool].setChecked(True)

        if tool == 'select':
            self.view.setDragMode(QGraphicsView.RubberBandDrag)
            self.view.setCursor(Qt.ArrowCursor)
        elif tool in ('text', 'edit_existing'):
            self.view.setDragMode(QGraphicsView.NoDrag)
            self.view.setCursor(Qt.IBeamCursor)
        else:
            self.view.setDragMode(QGraphicsView.NoDrag)
            self.view.setCursor(Qt.CrossCursor)

    def _pick_color(self):
        c = QColorDialog.getColor(self.current_color, self, "Цвет")
        if c.isValid():
            self.current_color = c
            self._update_color_btn()

    def _update_color_btn(self):
        self.color_btn.setStyleSheet(
            f"background:{self.current_color.name()};"
            f"border:1px solid #555; border-radius:6px;"
            f"min-height:{UI_ACTION_H - 2}px;"
        )
        self.color_btn.setText(self.current_color.name().upper())

    def _set_width(self, v):
        self.current_width = v
        self.width_lbl.setText(f"{v} px")

    # ─── Рендер ───
    def render_page(self):
        if not self.pdf or self.pdf.page_count == 0:
            return

        for items in self.annotations.values():
            for it in items:
                if it.scene() is self.scene:
                    self.scene.removeItem(it)

        if self.page_pixmap_item is not None:
            if self.page_pixmap_item.scene() is self.scene:
                self.scene.removeItem(self.page_pixmap_item)
            self.page_pixmap_item = None

        page = self.pdf[self.current_page]
        pix = page.get_pixmap(matrix=fitz.Matrix(ZOOM, ZOOM), alpha=False)
        img = pixmap_to_qimage(pix)

        self.page_pixmap_item = QGraphicsPixmapItem(QPixmap.fromImage(img))
        self.page_pixmap_item.setZValue(-100)
        self.scene.addItem(self.page_pixmap_item)

        for it in self.annotations.get(self.current_page, []):
            self.scene.addItem(it)

        self.scene.setSceneRect(0, 0, pix.width, pix.height)
        self.view.fitInView(self.scene.sceneRect(), Qt.KeepAspectRatio)

        self.page_lbl.setText(f"{self.current_page + 1} / {self.pdf.page_count}")

    def change_page(self, delta):
        if not self.pdf:
            return
        new_p = self.current_page + delta
        if 0 <= new_p < self.pdf.page_count:
            self.current_page = new_p
            self.render_page()

    # ─── Открытие / создание ───
    def open_pdf(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Открыть PDF", "", "PDF (*.pdf)")
        if not path:
            return
        try:
            doc = fitz.open(path)
            self.pdf = doc
            self.current_page = 0
            self.annotations = {}
            self.render_page()
            self.statusBar().showMessage(f"Открыт: {os.path.basename(path)}")
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", str(e))

    def new_document(self):
        doc = fitz.open()
        doc.new_page(width=595, height=842)
        self.pdf = doc
        self.current_page = 0
        self.annotations = {}
        self.render_page()
        self.statusBar().showMessage("Создан новый документ (A4)")

    def open_signatures_folder(self):
        try:
            if sys.platform.startswith("win"):
                os.startfile(SIGNATURES_DIR)
            elif sys.platform == "darwin":
                subprocess.Popen(["open", SIGNATURES_DIR])
            else:
                subprocess.Popen(["xdg-open", SIGNATURES_DIR])
        except Exception as e:
            QMessageBox.information(self, "Папка подписей",
                                    f"Путь:\n{SIGNATURES_DIR}\n\n{e}")

    # ─── Редактирование существующего текста со сдвигом ───
    def edit_existing_text_at(self, pos):
        if not self.pdf:
            return

        page = self.pdf[self.current_page]
        pdf_x = pos.x() * PDF_SCALE
        pdf_y = pos.y() * PDF_SCALE

        try:
            d = page.get_text("dict")
        except Exception as e:
            QMessageBox.warning(self, "Ошибка",
                                f"Не удалось прочитать текст:\n{e}")
            return

        # 1. Найти line под курсором
        target_line = None
        for block in d.get("blocks", []):
            if block.get("type", 0) != 0:
                continue
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    x0, y0, x1, y1 = span["bbox"]
                    if x0 <= pdf_x <= x1 and y0 <= pdf_y <= y1:
                        target_line = line
                        break
                if target_line:
                    break
            if target_line:
                break

        if target_line is None:
            QMessageBox.information(
                self, "Текст не найден",
                "Под курсором нет текстового фрагмента.\n\n"
                "• Кликните точнее по буквам\n"
                "• Если текст — картинка (скан),\n"
                "  редактирование невозможно"
            )
            return

        spans = target_line.get("spans", [])
        full_text = "".join(s.get("text", "") for s in spans)
        if not full_text.strip():
            return

        s0 = spans[0]
        size_pt = float(s0.get("size", 12))
        color_int = int(s0.get("color", 0))
        flags = int(s0.get("flags", 0))
        is_bold = bool(flags & 16)
        is_italic = bool(flags & 2)

        cr = (color_int >> 16) & 255
        cg = (color_int >> 8) & 255
        cb = color_int & 255

        lx0, ly0, lx1, ly1 = target_line["bbox"]

        style_bits = []
        if is_bold:
            style_bits.append("жирный")
        if is_italic:
            style_bits.append("курсив")
        style_str = "  •  " + ", ".join(style_bits) if style_bits else ""

        # 2. Диалог
        new_text, ok = QInputDialog.getMultiLineText(
            self, "Изменить текст",
            f"Редактирование (размер {size_pt:.1f} pt{style_str}):\n"
            f"Если строк станет больше — текст ниже\n"
            f"аккуратно сдвинется вниз. Размер шрифта не меняется.\n"
            f"(очистите поле, чтобы удалить строку)",
            full_text
        )
        if not ok:
            return

        # 3. Вычисляем сдвиг (в pt)
        line_height_pt = size_pt * LINE_HEIGHT_FACTOR
        if new_text.strip():
            n_lines_new = max(1, len(new_text.split("\n")))
        else:
            n_lines_new = 0
        n_lines_old = 1
        delta_pt = (n_lines_new - n_lines_old) * line_height_pt
        delta_scene = delta_pt * ZOOM

        # 4. Сдвигаем СУЩЕСТВУЮЩИЕ аннотации ниже линии
        if abs(delta_scene) > 0.01:
            shift_threshold_pt = ly1 - 0.5
            old_items = []
            for it in list(self.annotations.get(self.current_page, [])):
                it_y_pt = it.pos().y() * PDF_SCALE
                if it_y_pt >= shift_threshold_pt:
                    old_items.append(it)
            for it in old_items:
                it.setPos(it.pos().x(), it.pos().y() + delta_scene)

        # 5. Накрываем оригинальную строку белым
        pad = 1.5
        rx = (lx0 - pad) * ZOOM
        ry = (ly0 - pad) * ZOOM
        rw = ((lx1 - lx0) + pad * 2) * ZOOM
        rh = ((ly1 - ly0) + pad * 2) * ZOOM

        white = SelectableRectItem(QRectF(0, 0, rw, rh))
        white.setPos(rx, ry)
        white.setBrush(QBrush(QColor(255, 255, 255)))
        white.setPen(QPen(Qt.NoPen))
        white.setZValue(4)
        white.setFlag(QGraphicsItem.ItemIsMovable)
        white.setFlag(QGraphicsItem.ItemIsSelectable)
        white.setData(0, {'kind': 'rect', 'color': '#FFFFFF'})
        self.scene.addItem(white)
        self.register_annotation('rect', white)

        # 6. Рисуем новый текст (размер НЕ меняется)
        if new_text.strip():
            item = SelectableTextItem(new_text)
            item.setPos(lx0 * ZOOM, ly0 * ZOOM)

            font = QFont("Arial")
            font.setPixelSize(max(4, int(round(size_pt * ZOOM))))
            font.setBold(is_bold)
            font.setItalic(is_italic)
            item.setFont(font)

            item.setDefaultTextColor(QColor(cr, cg, cb))
            item.setTextInteractionFlags(Qt.NoTextInteraction)
            item.setFlag(QGraphicsItem.ItemIsMovable)
            item.setFlag(QGraphicsItem.ItemIsSelectable)
            item.setZValue(5)
            item.setData(0, {
                'kind': 'text',
                'color': f'#{cr:02x}{cg:02x}{cb:02x}',
                'bold': is_bold,
                'italic': is_italic,
                'pdf_size_pt': size_pt,
            })

            self.scene.addItem(item)
            self.register_annotation('text', item)

        # 7. Сдвигаем оригинальный PDF-текст ниже (перерисовываем)
        if abs(delta_pt) > 0.01:
            self._shift_pdf_spans_below(ly1, delta_pt)

        # Статус
        short = full_text[:40] + ('…' if len(full_text) > 40 else '')
        if not new_text.strip():
            self.statusBar().showMessage(f"Удалён текст «{short}»")
        elif abs(delta_pt) < 0.01:
            self.statusBar().showMessage(f"Заменён текст «{short}»")
        else:
            sign = "+" if delta_pt > 0 else ""
            self.statusBar().showMessage(
                f"Заменён текст «{short}» · сдвиг ниже {sign}{delta_pt:.1f} pt"
            )

    def _is_span_covered(self, span_bbox_pt):
        """Проверяет, покрыт ли span нашими белыми прямоугольниками."""
        sx0, sy0, sx1, sy1 = span_bbox_pt
        for item in self.annotations.get(self.current_page, []):
            data = item.data(0) or {}
            if data.get('kind') != 'rect':
                continue
            r = item.rect()
            p = item.pos()
            rx0 = (r.x() + p.x()) * PDF_SCALE
            ry0 = (r.y() + p.y()) * PDF_SCALE
            rx1 = (r.right() + p.x()) * PDF_SCALE
            ry1 = (r.bottom() + p.y()) * PDF_SCALE
            inter_x = max(0.0, min(sx1, rx1) - max(sx0, rx0))
            inter_y = max(0.0, min(sy1, ry1) - max(sy0, ry0))
            inter_area = inter_x * inter_y
            span_area = max(0.0001, (sx1 - sx0) * (sy1 - sy0))
            if inter_area / span_area > 0.5:
                return True
        return False

    def _shift_pdf_spans_below(self, y_threshold_pt, delta_pt):
        """Перерисовывает оригинальные span'ы PDF ниже y_threshold_pt со сдвигом."""
        page = self.pdf[self.current_page]

        try:
            d = page.get_text("dict")
        except Exception:
            return

        # Собираем span'ы ниже порога
        spans_below = []
        for block in d.get("blocks", []):
            if block.get("type", 0) != 0:
                continue
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    sx0, sy0, sx1, sy1 = span["bbox"]
                    if sy0 >= y_threshold_pt - 0.5:
                        text = span.get("text", "")
                        if not text.strip():
                            continue
                        if self._is_span_covered(span["bbox"]):
                            continue
                        spans_below.append(span)

        # Сортируем сверху вниз, чтобы z-order был стабильный
        spans_below.sort(key=lambda s: (s["bbox"][1], s["bbox"][0]))

        for s in spans_below:
            sx0, sy0, sx1, sy1 = s["bbox"]
            s_text = s.get("text", "").replace("\n", " ").strip()
            if not s_text:
                continue

            s_size_pt = float(s.get("size", 12))
            s_color_int = int(s.get("color", 0))
            s_flags = int(s.get("flags", 0))
            s_bold = bool(s_flags & 16)
            s_italic = bool(s_flags & 2)

            sc_r = (s_color_int >> 16) & 255
            sc_g = (s_color_int >> 8) & 255
            sc_b = s_color_int & 255

            # --- 1) Белый прямоугольник поверх оригинала ---
            pad = 0.5
            rx = (sx0 - pad) * ZOOM
            ry = (sy0 - pad) * ZOOM
            rw = ((sx1 - sx0) + pad * 2) * ZOOM
            rh = ((sy1 - sy0) + pad * 2) * ZOOM

            white = SelectableRectItem(QRectF(0, 0, rw, rh))
            white.setPos(rx, ry)
            white.setBrush(QBrush(QColor(255, 255, 255)))
            white.setPen(QPen(Qt.NoPen))
            white.setZValue(4)
            white.setFlag(QGraphicsItem.ItemIsMovable)
            white.setFlag(QGraphicsItem.ItemIsSelectable)
            white.setData(0, {'kind': 'rect', 'color': '#FFFFFF'})
            self.scene.addItem(white)
            self.register_annotation('rect', white)

            # --- 2) Новый QGraphicsTextItem со сдвигом ---
            new_y_pt = sy0 + delta_pt
            item = SelectableTextItem(s_text)
            item.setPos(sx0 * ZOOM, new_y_pt * ZOOM)

            font = QFont("Arial")
            font.setPixelSize(max(4, int(round(s_size_pt * ZOOM))))
            font.setBold(s_bold)
            font.setItalic(s_italic)
            item.setFont(font)

            item.setDefaultTextColor(QColor(sc_r, sc_g, sc_b))
            item.setTextInteractionFlags(Qt.NoTextInteraction)
            item.setFlag(QGraphicsItem.ItemIsMovable)
            item.setFlag(QGraphicsItem.ItemIsSelectable)
            item.setZValue(5)
            item.setData(0, {
                'kind': 'text',
                'color': f'#{sc_r:02x}{sc_g:02x}{sc_b:02x}',
                'bold': s_bold,
                'italic': s_italic,
                'pdf_size_pt': s_size_pt,
            })

            self.scene.addItem(item)
            self.register_annotation('text', item)

    # ─── Конвертация ───
    def convert_files(self):
        files, _ = QFileDialog.getOpenFileNames(
            self, "Файлы для конвертации", "",
            "Все поддерживаемые (*.png *.jpg *.jpeg *.bmp *.gif *.webp "
            "*.txt *.pdf *.docx *.doc *.pptx *.ppt *.xlsx *.xls *.odt *.odp *.ods);;"
            "Изображения (*.png *.jpg *.jpeg *.bmp *.gif *.webp);;"
            "Текст (*.txt);;PDF (*.pdf);;"
            "Word (*.docx *.doc);;"
            "PowerPoint (*.pptx *.ppt);;"
            "Excel (*.xlsx *.xls);;"
            "OpenDocument (*.odt *.odp *.ods);;"
            "Все файлы (*.*)"
        )
        if not files:
            return

        doc = fitz.open()
        errors = []

        for f in files:
            ext = os.path.splitext(f)[1].lower()
            try:
                if ext == '.pdf':
                    src = fitz.open(f)
                    doc.insert_pdf(src)
                    src.close()

                elif ext in ('.png', '.jpg', '.jpeg', '.bmp', '.gif', '.webp'):
                    page = doc.new_page(width=595, height=842)
                    rect = fitz.Rect(40, 40, 555, 802)
                    page.insert_image(rect, filename=f, keep_proportion=True)

                elif ext == '.txt':
                    with open(f, 'r', encoding='utf-8', errors='replace') as fh:
                        text = fh.read()
                    page = doc.new_page(width=595, height=842)
                    rect = fitz.Rect(50, 50, 545, 792)
                    kwargs = {}
                    if FONT_PATH:
                        kwargs['fontfile'] = FONT_PATH
                        kwargs['fontname'] = "custom"
                    else:
                        kwargs['fontname'] = "helv"
                    page.insert_textbox(rect, text, fontsize=12, **kwargs)

                elif ext in ('.docx', '.doc', '.pptx', '.ppt',
                             '.xlsx', '.xls', '.odt', '.odp', '.ods'):
                    pdf_tmp = libreoffice_convert(f)
                    if pdf_tmp:
                        src = fitz.open(pdf_tmp)
                        doc.insert_pdf(src)
                        src.close()
                        try:
                            os.remove(pdf_tmp)
                            shutil.rmtree(os.path.dirname(pdf_tmp),
                                          ignore_errors=True)
                        except Exception:
                            pass
                    else:
                        if ext in ('.docx', '.doc'):
                            docx_to_pdf_fallback(f, doc)
                        elif ext in ('.pptx', '.ppt'):
                            pptx_to_pdf_fallback(f, doc)
                        else:
                            errors.append(
                                f"{os.path.basename(f)} — для {ext} нужен "
                                f"LibreOffice"
                            )
                else:
                    errors.append(f"{os.path.basename(f)} — неизвестный формат")

            except Exception as e:
                errors.append(f"{os.path.basename(f)}: {e}")

        if doc.page_count == 0:
            msg = "Не удалось создать PDF."
            if errors:
                msg += "\n\n" + "\n".join(errors)
            QMessageBox.information(self, "Пусто", msg)
            return

        self.pdf = doc
        self.current_page = 0
        self.annotations = {}
        self.render_page()

        info = f"Создан PDF: {doc.page_count} стр."
        if not LIBREOFFICE:
            info += ("\n\n⚠ LibreOffice не найден — DOCX/PPTX/XLSX "
                     "конвертированы упрощённо (только текст).")
        if errors:
            info += "\n\nПредупреждения:\n" + "\n".join(errors)
        QMessageBox.information(self, "Готово", info)
        self.statusBar().showMessage(f"Конвертировано: {len(files)} файл(ов)")

    # ─── Аннотации ───
    def add_text_at(self, pos):
        text, ok = QInputDialog.getMultiLineText(self, "Добавить текст", "Текст:")
        if not ok or not text.strip():
            return

        item = SelectableTextItem(text)
        item.setPos(pos)
        font = QFont("Arial")
        font.setPixelSize(int(14 * ZOOM))
        item.setFont(font)
        item.setDefaultTextColor(self.current_color)
        item.setTextInteractionFlags(Qt.NoTextInteraction)
        item.setFlag(QGraphicsItem.ItemIsMovable)
        item.setFlag(QGraphicsItem.ItemIsSelectable)
        item.setZValue(5)
        item.setData(0, {'kind': 'text', 'color': self.current_color.name()})

        self.scene.addItem(item)
        self.register_annotation('text', item)

    def add_image_at(self, pos):
        path, _ = QFileDialog.getOpenFileName(
            self, "Изображение", "",
            "Изображения (*.png *.jpg *.jpeg *.bmp *.gif *.webp)"
        )
        if not path:
            return
        pm = QPixmap(path)
        if pm.isNull():
            return
        if pm.width() > 400:
            pm = pm.scaledToWidth(400, Qt.SmoothTransformation)

        item = SelectablePixmapItem(pm)
        item.setPos(pos)
        item.setFlag(QGraphicsItem.ItemIsMovable)
        item.setFlag(QGraphicsItem.ItemIsSelectable)
        item.setZValue(5)
        item.setData(0, {'kind': 'image', 'path': path})

        self.scene.addItem(item)
        self.register_annotation('image', item)

    def add_signature_at(self, pos):
        dlg = SignatureDialog(self)
        if dlg.exec() != QDialog.Accepted:
            return
        pm = dlg.result_pixmap
        if pm is None or pm.isNull():
            return

        item = SelectablePixmapItem(pm)
        item.setPos(pos)
        item.setFlag(QGraphicsItem.ItemIsMovable)
        item.setFlag(QGraphicsItem.ItemIsSelectable)
        item.setZValue(5)
        item.setData(0, {'kind': 'signature'})

        self.scene.addItem(item)
        self.register_annotation('signature', item)

    def register_annotation(self, kind, item):
        self.annotations.setdefault(self.current_page, []).append(item)

    def delete_selected(self):
        if self.current_page not in self.annotations:
            return
        items = list(self.annotations[self.current_page])
        for it in items:
            if it.isSelected():
                if it.scene() is self.scene:
                    self.scene.removeItem(it)
                self.annotations[self.current_page].remove(it)

    # ─── Экспорт ───
    def export_pdf(self):
        if not self.pdf:
            QMessageBox.information(self, "Нет документа",
                                    "Сначала откройте или создайте PDF.")
            return

        path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить PDF", "output.pdf", "PDF (*.pdf)"
        )
        if not path:
            return

        try:
            out = fitz.open()
            out.insert_pdf(self.pdf)

            for page_idx, items in self.annotations.items():
                if page_idx >= out.page_count:
                    continue
                page = out[page_idx]

                # ⚠ ВАЖНО: сортируем по zValue, чтобы белые прямоугольники
                # (z=4) наносились ДО текстов (z=5), иначе они затирают
                # уже нанесённый новый текст
                sorted_items = sorted(items, key=lambda it: it.zValue())

                for it in sorted_items:
                    self._apply_item(page, it)

            out.save(path, garbage=4, deflate=True)
            out.close()

            QMessageBox.information(self, "Готово", f"Сохранено:\n{path}")
            self.statusBar().showMessage(f"Сохранено: {os.path.basename(path)}")
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", str(e))

    def _apply_item(self, page, item):
        data = item.data(0) or {}
        kind = data.get('kind')

        if kind == 'text':
            text = item.toPlainText()
            pos = item.pos()

            if 'pdf_size_pt' in data:
                fs_pt = float(data['pdf_size_pt'])
            else:
                f = item.font()
                fs_px = f.pixelSize() if f.pixelSize() > 0 else max(f.pointSize(), 14)
                fs_pt = fs_px * PDF_SCALE

            is_bold = bool(data.get('bold', False))
            is_italic = bool(data.get('italic', False))
            if 'bold' not in data or 'italic' not in data:
                f = item.font()
                is_bold = f.bold()
                is_italic = f.italic()

            chosen_font = FONT_PATH
            if FONT_PATH and (is_bold or is_italic):
                variant = find_font_variant(FONT_PATH, is_bold, is_italic)
                if variant:
                    chosen_font = variant

            color = QColor(data.get('color', '#000000'))
            pdf_color = (color.redF(), color.greenF(), color.blueF())

            line_h = fs_pt * LINE_HEIGHT_FACTOR
            y_off = 0.0
            for line in text.split('\n'):
                bx = pos.x() * PDF_SCALE
                by = pos.y() * PDF_SCALE + fs_pt * 0.85 + y_off
                try:
                    if chosen_font:
                        page.insert_text(
                            fitz.Point(bx, by), line,
                            fontsize=fs_pt, fontfile=chosen_font,
                            fontname="custom", color=pdf_color
                        )
                    elif is_bold and is_italic:
                        page.insert_text(fitz.Point(bx, by), line,
                                         fontsize=fs_pt, fontname="hebi",
                                         color=pdf_color)
                    elif is_bold:
                        page.insert_text(fitz.Point(bx, by), line,
                                         fontsize=fs_pt, fontname="hebo",
                                         color=pdf_color)
                    elif is_italic:
                        page.insert_text(fitz.Point(bx, by), line,
                                         fontsize=fs_pt, fontname="heit",
                                         color=pdf_color)
                    else:
                        page.insert_text(fitz.Point(bx, by), line,
                                         fontsize=fs_pt, fontname="helv",
                                         color=pdf_color)
                except Exception as e:
                    print("text insert error:", e)
                y_off += line_h

        elif kind in ('image', 'signature'):
            rect = self._item_rect_in_pdf(item)
            try:
                if kind == 'image' and 'path' in data and os.path.exists(data['path']):
                    page.insert_image(rect, filename=data['path'],
                                      keep_proportion=True)
                else:
                    pm = item.pixmap()
                    if pm.isNull():
                        return
                    img = pm.toImage().convertToFormat(QImage.Format_ARGB32)
                    buf = QBuffer()
                    buf.open(QIODevice.WriteOnly)
                    img.save(buf, "PNG")
                    png = bytes(buf.data())
                    page.insert_image(rect, stream=png, keep_proportion=True)
            except Exception as e:
                print("image insert error:", e)

        elif kind in ('draw', 'whiteout'):
            color = QColor(data.get('color', '#000000'))
            pdf_color = (color.redF(), color.greenF(), color.blueF())
            width_px = data.get('width', 3)
            width_pt = width_px * PDF_SCALE

            path = item.path()
            offset = item.pos()
            pts = []
            for i in range(path.elementCount()):
                el = path.elementAt(i)
                pts.append(fitz.Point(
                    (el.x + offset.x()) * PDF_SCALE,
                    (el.y + offset.y()) * PDF_SCALE
                ))

            if len(pts) >= 2:
                shape = page.new_shape()
                for i in range(1, len(pts)):
                    shape.draw_line(pts[i - 1], pts[i])
                shape.finish(color=pdf_color, width=width_pt,
                             lineCap=1, lineJoin=1)
                shape.commit()

        elif kind == 'rect':
            color = QColor(data.get('color', '#FFFFFF'))
            pdf_color = (color.redF(), color.greenF(), color.blueF())
            r = item.rect()
            p = item.pos()
            rect_pdf = fitz.Rect(
                (r.x() + p.x()) * PDF_SCALE,
                (r.y() + p.y()) * PDF_SCALE,
                (r.right() + p.x()) * PDF_SCALE,
                (r.bottom() + p.y()) * PDF_SCALE,
            )
            shape = page.new_shape()
            shape.draw_rect(rect_pdf)
            shape.finish(color=None, fill=pdf_color)
            shape.commit()

    def _item_rect_in_pdf(self, item):
        r = item.sceneBoundingRect()
        return fitz.Rect(
            r.left() * PDF_SCALE,
            r.top() * PDF_SCALE,
            r.right() * PDF_SCALE,
            r.bottom() * PDF_SCALE,
        )


# ─────────────────────────────────────────────────────────────
#  Точка входа
# ─────────────────────────────────────────────────────────────
def main():
    QApplication.setHighDpiScaleFactorRoundingPolicy(
        Qt.HighDpiScaleFactorRoundingPolicy.PassThrough
    )

    app = QApplication(sys.argv)
    app.setStyle("Fusion")

    base = QFont("Segoe UI")
    base.setPointSize(UI_BASE_FONT)
    app.setFont(base)

    palette = app.palette()
    palette.setColor(QPalette.Window,           QColor("#1e1e24"))
    palette.setColor(QPalette.WindowText,       QColor("#eaeaea"))
    palette.setColor(QPalette.Base,             QColor("#2b2b33"))
    palette.setColor(QPalette.AlternateBase,    QColor("#26262e"))
    palette.setColor(QPalette.Text,             QColor("#eaeaea"))
    palette.setColor(QPalette.Button,           QColor("#383843"))
    palette.setColor(QPalette.ButtonText,       QColor("#eaeaea"))
    palette.setColor(QPalette.Highlight,        QColor("#ff6b6b"))
    palette.setColor(QPalette.HighlightedText,  QColor("#ffffff"))
    palette.setColor(QPalette.ToolTipBase,      QColor("#2b2b33"))
    palette.setColor(QPalette.ToolTipText,      QColor("#eaeaea"))
    palette.setColor(QPalette.PlaceholderText,  QColor("#888888"))
    app.setPalette(palette)

    win = MainWindow()
    win.showMaximized()
    sys.exit(app.exec())


if __name__ == '__main__':
    main()